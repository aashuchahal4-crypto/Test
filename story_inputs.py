import html
import os
import re
import zipfile
import xml.etree.ElementTree as ET

from script_generator import parse_custom_script
from video_generator import split_text_into_segments

SPEAKER_LINE = re.compile(r"^\s*[^:\n]{1,40}\s*:\s*\S+")
SENTENCE_RE = re.compile(r"(?<=[.!?।])\s+")


def extract_docx_text(path):
    if not path:
        return ""
    try:
        from docx import Document
        doc = Document(path)
        parts = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except ImportError:
        return _extract_docx_text_fallback(path)
    except Exception:
        return _extract_docx_text_fallback(path)


def _extract_docx_text_fallback(path):
    try:
        with zipfile.ZipFile(path) as zf:
            xml = zf.read("word/document.xml")
        root = ET.fromstring(xml)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        paragraphs = []
        for para in root.findall(".//w:p", ns):
            texts = [node.text or "" for node in para.findall(".//w:t", ns)]
            value = html.unescape("".join(texts)).strip()
            if value:
                paragraphs.append(value)
        return "\n".join(paragraphs).strip()
    except Exception as exc:
        raise ValueError(f"Could not extract DOCX text: {exc}") from exc


def combine_inputs(docx_path=None, raw_text=""):
    parts = []
    notes = []
    docx_text = extract_docx_text(docx_path) if docx_path else ""
    raw_text = (raw_text or "").strip()
    if docx_path:
        if docx_text:
            parts.append(docx_text)
            notes.append("DOCX text extracted first")
        else:
            notes.append("DOCX contained no readable text")
    if raw_text:
        parts.append(raw_text)
        notes.append("raw text appended")
    combined = "\n\n".join(parts).strip()
    if not combined:
        raise ValueError("No story text found. Upload a DOCX with readable paragraphs or paste text.")
    return combined, "; ".join(notes) if notes else "raw text used"


def looks_like_custom_script(text):
    lines = [line for line in (text or "").splitlines() if line.strip()]
    if not lines:
        return False
    matches = sum(1 for line in lines if SPEAKER_LINE.match(line))
    return matches >= max(1, min(2, len(lines))) or matches / len(lines) >= 0.4


def segment_story(text, max_chars=110, custom_script_mode="Auto"):
    text = (text or "").strip()
    if not text:
        raise ValueError("Story text is empty after extraction.")
    use_script = custom_script_mode == "Custom script" or (custom_script_mode == "Auto" and looks_like_custom_script(text))
    if use_script:
        dialogues = parse_custom_script(text)
        return _split_dialogues(dialogues, max_chars), "speaker-prefixed custom script"
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    scenes = []
    for para in paragraphs:
        sentences = [s.strip() for s in SENTENCE_RE.split(para) if s.strip()]
        buffer = ""
        for sent in sentences or [para]:
            if len(buffer) + len(sent) + 1 <= max_chars:
                buffer = (buffer + " " + sent).strip()
            else:
                if buffer:
                    scenes.extend(split_text_into_segments(buffer, max_chars=max_chars))
                buffer = sent
        if buffer:
            scenes.extend(split_text_into_segments(buffer, max_chars=max_chars))
    scenes = [s.strip() for s in scenes if s.strip()]
    if not scenes:
        raise ValueError("Could not segment story into scenes.")
    return [{"speaker": "NARRATOR", "text": scene, "emotion": "neutral"} for scene in scenes], "narrator line-by-line scenes"


def _split_dialogues(dialogues, max_chars):
    output = []
    for d in dialogues:
        chunks = split_text_into_segments(d.get("text", ""), max_chars=max_chars)
        for chunk in chunks:
            if chunk.strip():
                output.append({"speaker": d.get("speaker", "NARRATOR"), "text": chunk.strip(), "emotion": d.get("emotion", "neutral")})
    return output
